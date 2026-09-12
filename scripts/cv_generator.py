#!/usr/bin/env python3
"""
CV Generator for User Surname — Machine-Readable DOCX
Implements §0.4 Reference-CV Formatting Lock from CV_REPOSITORY_DATABASE.md v26.

Usage:
    python3 cv_generator.py <json_file>

Where <json_file> contains the tailored CV content as structured JSON.
See README.md for the expected JSON schema.

Output: User_CV_<Company>_<Role>.docx in the same directory.
"""

import json
import sys
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# §0.4 FORMATTING CONSTANTS
COLOR_NAME = RGBColor(0x36, 0x5F, 0x91)
COLOR_NAME_RULE = RGBColor(0x4F, 0x81, 0xBD)
COLOR_HEADING = RGBColor(0x1F, 0x49, 0x7D)
COLOR_BODY = RGBColor(0x00, 0x00, 0x00)
COLOR_FOOTER = RGBColor(0x00, 0x00, 0x00)
FONT_FAMILY = "Arial"

PAGE_WIDTH = Inches(8.2677)
PAGE_HEIGHT = Inches(11.6929)
MARGIN_TB = Inches(0.45)
MARGIN_LR = Inches(0.6201)


def add_bottom_border(paragraph, color_hex="4F81BD", size="8", space="4"):
    pPr = paragraph._element.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), size)
    bottom.set(qn('w:space'), space)
    bottom.set(qn('w:color'), color_hex)
    pbdr.append(bottom)
    pPr.append(pbdr)


def set_character_spacing(paragraph, points):
    pPr = paragraph._element.get_or_add_pPr()
    rPr = pPr.find(qn('w:rPr'))
    if rPr is None:
        rPr = OxmlElement('w:rPr')
        pPr.append(rPr)
    spacing = OxmlElement('w:spacing')
    spacing.set(qn('w:val'), str(int(points * 20)))
    rPr.append(spacing)


def set_keep_with_next(paragraph):
    pPr = paragraph._element.get_or_add_pPr()
    pPr.append(OxmlElement('w:keepNext'))
    pPr.append(OxmlElement('w:keepLines'))


def add_manual_line_break(paragraph):
    run = paragraph.add_run()
    br = OxmlElement('w:br')
    run._element.append(br)


def set_hanging_indent(paragraph, left, hanging):
    pPr = paragraph._element.get_or_add_pPr()
    ind = OxmlElement('w:ind')
    ind.set(qn('w:left'), str(int(left.inches * 1440)))
    ind.set(qn('w:hanging'), str(int(hanging.inches * 1440)))
    pPr.append(ind)


def add_section_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    set_keep_with_next(p)
    run = p.add_run(text.upper())
    run.font.name = FONT_FAMILY
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = COLOR_HEADING
    add_bottom_border(p, color_hex="1F497D", size="8", space="1")


def add_bullet(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    set_hanging_indent(p, Inches(0.18), Inches(0.12))
    run = p.add_run("\u2022 " + text)
    run.font.name = FONT_FAMILY
    run.font.size = Pt(10)
    run.font.color.rgb = COLOR_BODY


def build_cv(content, output_path):
    doc = Document()
    section = doc.sections[0]
    section.page_width = PAGE_WIDTH
    section.page_height = PAGE_HEIGHT
    section.top_margin = MARGIN_TB
    section.bottom_margin = MARGIN_TB
    section.left_margin = MARGIN_LR
    section.right_margin = MARGIN_LR

    normal = doc.styles['Normal']
    normal.font.name = FONT_FAMILY
    normal.font.size = Pt(10)
    normal.font.color.rgb = COLOR_BODY
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1.0

    # 1. Name — 20pt bold #365F91, centered, bottom rule
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run("User MARKOVI\u0106")
    run.font.name = FONT_FAMILY
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = COLOR_NAME
    set_character_spacing(p, 0.25)
    add_bottom_border(p, "4F81BD", "8", "4")

    # 2. Headline — 12pt bold italic, centered
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(content['headline'].upper())
    run.font.name = FONT_FAMILY
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.italic = True
    set_character_spacing(p, 0.75)

    # 3. Contact — 9pt, centered, two lines
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    r1 = p.add_run("Phone +381 XX XXX XXXX | email : your-email@example.com  | linkedin.com/in/user-handle)
    r1.font.name = FONT_FAMILY; r1.font.size = Pt(9); r1.font.color.rgb = COLOR_BODY
    add_manual_line_break(p)
    r2 = p.add_run("Belgrade, Serbia | EU citizen (Czech Republic)")
    r2.font.name = FONT_FAMILY; r2.font.size = Pt(9); r2.font.color.rgb = COLOR_BODY

    # Profile
    add_section_heading(doc, "PROFESSIONAL PROFILE")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(content['profile'])
    run.font.name = FONT_FAMILY; run.font.size = Pt(10); run.font.color.rgb = COLOR_BODY

    # Competencies
    add_section_heading(doc, "CORE COMPETENCIES")
    for comp in content['competencies']:
        add_bullet(doc, comp)

    # Experience
    add_section_heading(doc, "PROFESSIONAL EXPERIENCE")
    for exp in content['experience']:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        set_keep_with_next(p)
        role_text = exp['title']
        if exp.get('org'): role_text += " | " + exp['org']
        if exp.get('location'): role_text += " | " + exp['location']
        run = p.add_run(role_text)
        run.font.name = FONT_FAMILY; run.font.size = Pt(10); run.font.bold = True; run.font.color.rgb = COLOR_BODY

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        set_keep_with_next(p)
        run = p.add_run(exp['dates'])
        run.font.name = FONT_FAMILY; run.font.size = Pt(10); run.font.color.rgb = COLOR_BODY

        for bullet in exp['bullets']:
            add_bullet(doc, bullet)

    # Additional Experience
    if content.get('additional_experience'):
        add_section_heading(doc, "ADDITIONAL RELEVANT EXPERIENCE")
        for line in content['additional_experience']:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.0
            run = p.add_run(line)
            run.font.name = FONT_FAMILY; run.font.size = Pt(10); run.font.color.rgb = COLOR_BODY

    # Education
    add_section_heading(doc, "EDUCATION")
    for edu in content['education']:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.left_indent = Inches(0.5)
        run = p.add_run(edu)
        run.font.name = FONT_FAMILY; run.font.size = Pt(10); run.font.color.rgb = COLOR_BODY

    # Languages
    add_section_heading(doc, "LANGUAGES")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(content['languages'])
    run.font.name = FONT_FAMILY; run.font.size = Pt(10); run.font.color.rgb = COLOR_BODY

    # Footer
    footer = section.footer
    footer.is_linked_to_previous = False
    pf = footer.paragraphs[0]
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf.paragraph_format.line_spacing = 1.0
    run = pf.add_run("User Markovi\u0107")
    run.font.name = FONT_FAMILY; run.font.size = Pt(9); run.font.color.rgb = COLOR_FOOTER

    doc.save(output_path)
    return output_path


def main():
    if len(sys.argv) < 2:
        print("Usage: python3 cv_generator.py <json_file>")
        sys.exit(1)
    json_path = sys.argv[1]
    with open(json_path, 'r', encoding='utf-8') as f:
        content = json.load(f)
    company = content.get('company', 'Unknown')
    role = content.get('role', 'Role')
    output_dir = os.path.dirname(os.path.abspath(json_path))
    output_path = os.path.join(output_dir, f"User_CV_{company}_{role}.docx")
    build_cv(content, output_path)
    print(f"OK: {output_path}")


if __name__ == "__main__":
    main()
